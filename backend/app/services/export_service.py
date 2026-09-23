from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.config import Settings
from app.utils.errors import PipelineError


def timestamp(seconds: float) -> str:
    value = int(seconds)
    return f"{value // 3600:02}:{value // 60 % 60:02}:{value % 60:02}"


def sections(snapshot: dict):
    yield (
        "Участники",
        [
            p["name"] + (f" — {p['position']}" if p["position"] else "")
            for p in snapshot["participants"]
        ]
        or ["Не указаны"],
    )
    yield "Краткое содержание", snapshot["summary"].split("\n\n") or ["Речь не обнаружена."]
    yield "Темы", snapshot["topics"] or ["Не выделены"]
    yield "Решения", snapshot["decisions"] or ["Явные решения не обнаружены"]
    assignments = []
    for index, task in enumerate(snapshot["tasks"], 1):
        deadline = task["deadline"] or "не указан"
        assignments.append(f"{index}. {task['description']}")
        assignments.append(
            f"Ответственный: {task['responsible_name'] or 'не указан'} | Автор: {task['assigned_by'] or 'не указан'}"
        )
        assignments.append(
            f"Срок: {deadline} | Статус: {task['status']} | Приоритет: {task['priority']} | Уверенность: {task['confidence']:.0%}"
        )
        assignments.append(f"Основание: {task['source_quote']}")
        if task["analysis_version"] < snapshot["analysis_version"]:
            assignments.append("Сохранено из предыдущего анализа; проверьте актуальность.")
    yield "Поручения", assignments or ["Явные поручения не обнаружены"]
    yield (
        "Транскрипт",
        [
            f"[{timestamp(s['start'])}–{timestamp(s['end'])}] {s['speaker_name'] or s['speaker_id']}: {s['text']}"
            for s in snapshot["transcript"]
        ]
        or ["Речь не обнаружена"],
    )


class ExportService:
    def __init__(self, settings: Settings):
        self.settings = settings

    def docx(self, snapshot: dict) -> bytes:
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2.2)
        normal = document.styles["Normal"]
        normal.font.name, normal.font.size = "DejaVu Sans", Pt(10)
        normal.paragraph_format.space_after = Pt(6)
        for name in ("Title", "Heading 1"):
            document.styles[name].font.name = "DejaVu Sans"
            document.styles[name].font.color.rgb = RGBColor.from_string("173B56")
        document.add_heading("ПРОТОКОЛ СОВЕЩАНИЯ", 0)
        document.add_heading(snapshot["title"], 1)
        document.add_paragraph(f"Дата: {snapshot['meeting_date']} ({snapshot['timezone']})")
        document.add_paragraph(
            "Автоматически подготовленный проект. Проверьте имена, решения и поручения по записи."
        )
        for heading, paragraphs in sections(snapshot):
            document.add_heading(heading, 1)
            for paragraph in paragraphs:
                document.add_paragraph(paragraph)
        section.footer.paragraphs[0].text = "Протокол • Локальная обработка"
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def pdf(self, snapshot: dict) -> bytes:
        for path in (self.settings.pdf_font_path, self.settings.pdf_bold_font_path):
            if not Path(path).is_file():
                raise PipelineError("PDF_FONT_MISSING")
        if "MeetingSans" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("MeetingSans", str(self.settings.pdf_font_path)))
            pdfmetrics.registerFont(
                TTFont("MeetingSans-Bold", str(self.settings.pdf_bold_font_path))
            )
            pdfmetrics.registerFontFamily(
                "MeetingSans",
                normal="MeetingSans",
                bold="MeetingSans-Bold",
                italic="MeetingSans",
                boldItalic="MeetingSans-Bold",
            )
        styles = getSampleStyleSheet()
        styles.add(
            ParagraphStyle("BodyRU", fontName="MeetingSans", fontSize=10, leading=15, spaceAfter=7)
        )
        styles.add(
            ParagraphStyle(
                "HeadingRU",
                fontName="MeetingSans-Bold",
                fontSize=14,
                leading=19,
                spaceBefore=16,
                spaceAfter=8,
                keepWithNext=True,
                textColor=colors.HexColor("#173B56"),
            )
        )
        styles.add(
            ParagraphStyle(
                "TitleRU", parent=styles["HeadingRU"], fontSize=20, leading=26, alignment=TA_CENTER
            )
        )

        def paragraph(text, style="BodyRU"):
            return Paragraph(escape(str(text)).replace("\n", "<br/>"), styles[style])

        story = [
            paragraph("ПРОТОКОЛ СОВЕЩАНИЯ", "TitleRU"),
            paragraph(snapshot["title"], "HeadingRU"),
            paragraph(f"Дата: {snapshot['meeting_date']} ({snapshot['timezone']})"),
            paragraph(
                "Автоматически подготовленный проект. Проверьте имена, решения и поручения по записи."
            ),
            Spacer(1, 0.2 * cm),
        ]
        for heading, paragraphs in sections(snapshot):
            story.append(paragraph(heading, "HeadingRU"))
            story.extend(paragraph(text) for text in paragraphs)

        def footer(canvas, document):
            canvas.saveState()
            canvas.setFont("MeetingSans", 8)
            canvas.drawString(2 * cm, 1.2 * cm, "Протокол • Локальная обработка")
            canvas.drawRightString(document.pagesize[0] - 2 * cm, 1.2 * cm, str(document.page))
            canvas.restoreState()

        buffer = BytesIO()
        SimpleDocTemplate(
            buffer,
            leftMargin=2 * cm,
            rightMargin=2 * cm,
            topMargin=1.8 * cm,
            bottomMargin=2 * cm,
            title="Протокол совещания",
            author="Local Meeting Backend",
        ).build(story, onFirstPage=footer, onLaterPages=footer)
        return buffer.getvalue()
